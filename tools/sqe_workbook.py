#!/usr/bin/env python3
"""SQE A2 — build the testing workbook (.xlsx, exactly 2 sheets) from CSV sources + gtest XML (schema: docs/specs/SPEC_05).
Usage:
  python3 tools/sqe_workbook.py --inventory work/inventory/test_inventory.csv --mcdc work/mcdc/mcdc_matrix.csv \
      --xml evidence/tests/xml evidence/tests/probes [--out deliverables/<BASE>.xlsx] [--validate-only]
Execution Result is taken from XML (never typed by hand); rows without XML need manual_status BLOCKED/NOT EXECUTED + manual_note."""
import argparse, csv, os, subprocess, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sqe_trace_check import load_results
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

INV_COLS = [("Test ID", 'test_id', 16, "SQE-… identifier (SPEC_01); encoded in the gtest name"),
            ("Component / Function", 'component', 28, "Production class::function under test"),
            ("Purpose / Scenario", 'purpose', 44, "One sentence; characterization/probe tests say so"),
            ("Key Controlled Input / State", 'controlled_input', 40, "Inputs, params, uORB messages, prior calls, time"),
            ("Expected Result", 'expected', 46, "Oracle with its source in brackets"),
            ("Execution Result", None, 16, "From gtest XML: PASS / FAIL / BLOCKED / NOT EXECUTED"),
            ("Structural Coverage Target", 'coverage_target', 30, "Decision IDs + outcomes (REF_03–REF_05)"),
            ("Test File / Evidence Reference", None, 52, "Source file::test and XML evidence path")]
MC_COLS = [("Decision ID", 'decision_id', 11), ("Source", 'source', 24), ("Decision", 'decision', 28), ("Conditions", 'conditions', 46),
           ("Test ID", 'test_id', 15), ("Evaluation", 'evaluation', 11)] + [(f"C{k}", f"c{k}", 7) for k in range(1, 8)] + \
          [("Outcome", 'outcome', 9), ("Independence Pair", 'pair_with', 30), ("Condition Shown Independent", 'shows', 13),
           ("MC/DC Form", 'form', 13), ("Observability", 'observability', 18), ("Notes", 'notes', 40)]
FILL = {k: PatternFill('solid', start_color=v, end_color=v) for k, v in
        {'hdr': '1F5F6B', 'PASS': 'C6EFCE', 'FAIL': 'FFC7CE', 'BLOCKED': 'FFEB9C', 'NOT': 'E7E6E6', 'T': 'E2F0D9', 'F': 'FCE4D6', 'NE': 'EDEDED'}.items()}
THIN = Border(*(Side(style='thin', color='BFBFBF'),) * 4)

def base_name(root):
    try:
        for line in open(os.path.join(root, 'work', 'TEAM.md')):
            if line.startswith('BASE='): return line.strip()[5:]
    except OSError: pass
    return 'SQE_A2'

def header(ws, cols, notes=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=c[0]); cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = FILL['hdr']; cell.alignment = Alignment(wrap_text=True, vertical='center'); cell.border = THIN
        ws.column_dimensions[get_column_letter(i)].width = c[2]
        if notes and notes[i-1]: cell.comment = Comment(notes[i-1], 'SQE A2')
    ws.freeze_panes = 'A2'; ws.row_dimensions[1].height = 30

def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--inventory', required=True); ap.add_argument('--mcdc', required=True)
    ap.add_argument('--xml', nargs='*', default=['evidence/tests/xml', 'evidence/tests/probes']); ap.add_argument('--out')
    ap.add_argument('--validate-only', action='store_true'); a = ap.parse_args()
    root = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
    inv = list(csv.DictReader(open(a.inventory, newline='', encoding='utf-8')))
    mc = list(csv.DictReader(open(a.mcdc, newline='', encoding='utf-8')))
    res = load_results([p for p in a.xml if os.path.exists(p)])
    errors, ids = [], [r['test_id'].strip() for r in inv]
    for d in sorted({i for i in ids if ids.count(i) > 1}): errors.append(f"duplicate Test ID {d}")
    rows = []
    for r in inv:
        for k in ('test_id', 'component', 'purpose', 'controlled_input', 'expected', 'coverage_target', 'gtest_name'):
            if not (r.get(k) or '').strip(): errors.append(f"{r.get('test_id')}: empty '{k}'")
        g = (r.get('gtest_name') or '').strip()
        if g in res: status, ev = res[g]
        elif (r.get('manual_status') or '').strip() in ('BLOCKED', 'NOT EXECUTED'):
            status, ev = f"{r['manual_status'].strip()} — {r.get('manual_note','').strip()}", 'no XML (manual status)'
            if not (r.get('manual_note') or '').strip(): errors.append(f"{r['test_id']}: manual status needs manual_note")
        else:
            status, ev = None, None; errors.append(f"{r['test_id']}: no XML result for '{g}' and no manual status")
        rows.append((r, status, ev))
    idset = set(ids)
    for r in mc:
        if r['test_id'].strip() not in idset: errors.append(f"MC/DC row {r['decision_id']} {r['evaluation']}: unknown test {r['test_id']}")
    chk = subprocess.run([sys.executable, os.path.join(os.path.dirname(__file__), 'sqe_mcdc_check.py'), a.mcdc], capture_output=True, text=True)
    if chk.returncode != 0: errors.append("MC/DC matrix incomplete — run tools/sqe_mcdc_check.py"); print(chk.stdout[-1500:])
    if errors:
        print("VALIDATION FAILED:"); [print("  -", e) for e in errors[:60]]; print(f"  ({len(errors)} errors)"); return 1
    print(f"validation OK: {len(rows)} inventory rows, {len(mc)} MC/DC rows, {sum(1 for _, s, _ in rows if s == 'PASS')} PASS")
    if a.validate_only: return 0
    wb = Workbook(); ws = wb.active; ws.title = 'Test Inventory'
    header(ws, INV_COLS, [c[3] for c in INV_COLS])
    for i, (r, status, ev) in enumerate(rows, 2):
        vals = [r['test_id'], r['component'], r['purpose'], r['controlled_input'], r['expected'], status, r['coverage_target'],
                f"{r.get('test_ref','').strip()} ; {ev}"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v); c.font = Font(name='Arial', size=10); c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = THIN
        key = 'PASS' if status == 'PASS' else 'FAIL' if status == 'FAIL' else 'BLOCKED' if str(status).startswith('BLOCKED') else 'NOT'
        ws.cell(row=i, column=6).fill = FILL[key]
    ws.auto_filter.ref = f"A1:{get_column_letter(len(INV_COLS))}{len(rows)+1}"
    ws2 = wb.create_sheet('MC-DC Evidence')
    notes = [None] * len(MC_COLS)
    notes[6] = "Condition values per evaluation: T/F = evaluated; NE(T)/NE(F) = not evaluated (short-circuit), logical value in brackets. Legend per decision in 'Conditions'."
    notes[14] = "Partner row(s) of the independence pair (TESTID or TESTID@evaluation), aligned with 'Condition Shown Independent'."
    notes[17] = "O1 return value, O2 state getter, O3 per-test structural evidence (path)"
    header(ws2, MC_COLS, notes)
    for i, r in enumerate(mc, 2):
        for j, (_, k, _) in enumerate(MC_COLS, 1):
            v = (r.get(k) or '').strip(); c = ws2.cell(row=i, column=j, value=v); c.font = Font(name='Arial', size=10)
            c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = THIN
            if k.startswith('c') and k[1:].isdigit() and v: c.fill = FILL['NE' if v.startswith('NE') else v[0]]
            if k == 'outcome' and v in ('T', 'F'): c.fill = FILL[v]
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(MC_COLS))}{len(mc)+1}"
    out = a.out or os.path.join(root, 'deliverables', base_name(root) + '.xlsx'); os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    if len(wb.sheetnames) != 2: print("ERROR: workbook does not have exactly 2 sheets"); return 1
    print(f"wrote {out} (sheets: {wb.sheetnames})"); return 0

if __name__ == '__main__':
    sys.exit(main())
